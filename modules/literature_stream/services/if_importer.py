"""
IF导入与回填工具：从Excel读取 SCIE 期刊的 ISSN 与 5年IF，并写入 journal_metadata，
随后可用于新文献入库时的即时匹配，或对近期缺失IF的文献做回填。

数据来源：通过环境变量 IF_EXCEL_PATH 指定 Excel 路径。
依赖：openpyxl（避免新增大量依赖）
"""
from __future__ import annotations

import os
from typing import Optional, Tuple
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text

def _norm_issn(raw: Optional[str]) -> Optional[str]:
    if not raw:
        return None
    s = str(raw).strip().upper().replace(" ", "")
    s = s.replace("ISSN", "").replace(":", "")
    s = s.replace("-", "")
    if len(s) == 8 and s.isalnum():
        return f"{s[:4]}-{s[4:]}"
    return None


def _guess_header_indexes(header_row) -> Tuple[Optional[int], Optional[int], Optional[int], Optional[int], Optional[int], Optional[int]]:
    """
    自动识别表头中 ISSN、E-ISSN、5年IF、中科院分区、索引库、期刊全名 列的索引（0-based）。
    返回: (issn_idx, eissn_idx, if5_idx, cas_idx, indexlib_idx, name_idx)
    """
    issn_idx = eissn_idx = if5_idx = cas_idx = indexlib_idx = name_idx = None
    for i, cell in enumerate(header_row):
        val = str(cell.value).strip() if cell.value is not None else ""
        v = val.lower()
        if issn_idx is None and ("issn" in v) and ("e-issn" not in v and "eissn" not in v and "electronic" not in v):
            issn_idx = i
        if eissn_idx is None and any(k in v for k in ["e-issn", "eissn", "electronic issn", "在线", "电子"]):
            eissn_idx = i
        # 5年IF 列：兼容 "5年IF"、"5 年IF" 等各种写法
        if if5_idx is None:
            vs = v.replace(" ", "")  # 去掉空格提高兼容性
            if any(k in vs for k in ["5年if", "五年if", "5yearif", "5-yearif", "5yrif"]):
                if5_idx = i
        # 中科院分区（2025）列
        if cas_idx is None and any(k in v for k in ["中科院分区", "cas zone", "2025分区"]):
            cas_idx = i
        if indexlib_idx is None and any(k in v for k in ["索引库", "index", "数据库", "collection"]):
            indexlib_idx = i
        if name_idx is None and any(k in v for k in ["期刊", "全名", "刊名", "journal", "title", "full", "名称"]):
            name_idx = i
    return issn_idx, eissn_idx, if5_idx, cas_idx, indexlib_idx, name_idx


async def import_if_from_excel(db: AsyncSession) -> int:
    """
    从 Excel 导入/更新 journal_metadata 的 if_5y 与分区。
    返回导入条数。
    """
    path = os.getenv("IF_EXCEL_PATH")
    if not path or not os.path.exists(path):
        # 没有文件则跳过
        return 0

    wb = load_workbook(filename=path, read_only=True, data_only=True)
    total_upserts = 0

    async with db.begin():
        # 建立临时表（注意 asyncpg 不允许一条语句中包含多个命令）
        # 现在同时携带 full_name，方便后续与 JCR 表按刊名对齐
        await db.execute(text("""
            CREATE TEMP TABLE IF NOT EXISTS _if_import(
              issn varchar(20),
              full_name varchar(300),
              if_5y numeric(5,2),
              cas_zone varchar(10)
            );
        """))
        await db.execute(text("TRUNCATE _if_import;"))

        for sheet in wb.worksheets:
            # Sheet名称包含SCIE，则认为整个sheet都是SCIE期刊
            is_scie_sheet = "SCIE" in sheet.title.upper()
            
            rows = sheet.iter_rows(values_only=False)
            try:
                header = next(rows)
            except StopIteration:
                continue
            issn_idx, eissn_idx, if5_idx, cas_idx, idxlib_idx, name_idx = _guess_header_indexes(header)
            if (issn_idx is None and eissn_idx is None) or if5_idx is None:
                continue

            for row in rows:
                issn_raw = row[issn_idx].value if (issn_idx is not None and issn_idx < len(row)) else None
                eissn_raw = row[eissn_idx].value if (eissn_idx is not None and eissn_idx < len(row)) else None
                if5_raw = row[if5_idx].value if if5_idx < len(row) else None
                name_raw = row[name_idx].value if (name_idx is not None and name_idx < len(row)) else None
                cas_raw = row[cas_idx].value if (cas_idx is not None and cas_idx < len(row)) else None
                
                # 仅保留 SCIE：如果是SCIE sheet则跳过此检查，否则检查索引库列
                if not is_scie_sheet and idxlib_idx is not None:
                    idx_val = row[idxlib_idx].value
                    if idx_val and "scie" not in str(idx_val).lower():
                        continue
                # 可能同时存在 ISSN 和 E-ISSN：都入
                norm_list = []
                p_issn = _norm_issn(str(issn_raw) if issn_raw is not None else None)
                e_issn = _norm_issn(str(eissn_raw) if eissn_raw is not None else None)
                if p_issn:
                    norm_list.append(p_issn)
                if e_issn and e_issn != p_issn:
                    norm_list.append(e_issn)
                if not norm_list:
                    continue
                try:
                    if5 = float(if5_raw) if if5_raw is not None and str(if5_raw).strip() != "" else None
                except Exception:
                    continue
                fullname = str(name_raw).strip() if name_raw else None
                cas_zone = None
                if cas_raw is not None:
                    cas_zone = str(cas_raw).strip()
                    if cas_zone == "":
                        cas_zone = None
                for issn in norm_list:
                    await db.execute(
                        text("INSERT INTO _if_import(issn, full_name, if_5y, cas_zone) VALUES (:issn, :fullname, :if5, :cas_zone)"),
                        {"issn": issn, "fullname": fullname, "if5": if5, "cas_zone": cas_zone}
                    )
                    # 兼容旧数据：若 journal_metadata 中已有该 ISSN 且 full_name 为空/占位，则用当前的刊名回填
                    if fullname:
                        await db.execute(
                            text("""
                                UPDATE journal_metadata SET full_name = :fullname, updated_at = NOW()
                                WHERE issn = :issn AND (full_name IS NULL OR full_name = issn)
                            """),
                            {"fullname": fullname, "issn": issn}
                        )

        # upsert 到 journal_metadata
        # 注意：_if_import 可能包含同一 ISSN 的多行，必须先按 ISSN 去重，
        # 否则单条 INSERT ... ON CONFLICT 语句会多次命中同一行，触发
        # "ON CONFLICT DO UPDATE command cannot affect row a second time" 错误。
        # 同时在去重阶段携带 full_name，确保后续可以按刊名与 JCR 缩写表对齐。
        res = await db.execute(text("""
            WITH dedup AS (
              SELECT
                issn,
                MAX(full_name) AS full_name,
                MAX(if_5y) AS if_5y,
                MAX(cas_zone) AS cas_zone
              FROM _if_import
              WHERE issn IS NOT NULL
              GROUP BY issn
            ), up AS (
              INSERT INTO journal_metadata(issn, if_5y, cas_zone, full_name)
              SELECT d.issn,
                     d.if_5y,
                     COALESCE(d.cas_zone, jm.cas_zone),
                     COALESCE(jm.full_name, d.full_name)
              FROM dedup d
              LEFT JOIN journal_metadata jm ON jm.issn = d.issn
              ON CONFLICT (issn) DO UPDATE
              SET if_5y = EXCLUDED.if_5y,
                  cas_zone = COALESCE(EXCLUDED.cas_zone, journal_metadata.cas_zone),
                  full_name = COALESCE(journal_metadata.full_name, EXCLUDED.full_name),
                  updated_at = NOW()
              RETURNING 1
            )
            SELECT COUNT(*) FROM up;
        """))
        total_upserts = int(res.scalar() or 0)

        # =============
        # JCR 缩写导入（期刊名称缩写 -> 期刊全名）
        # 使用同一个 Excel 中的 JCR sheet，将缩写写入 journal_metadata.nlm_abbr，
        # 便于后续使用 PubMed 的期刊缩写做 IF 映射。
        # =============
        for sheet in wb.worksheets:
            if "JCR" not in sheet.title.upper():
                continue

            rows = sheet.iter_rows(values_only=False)
            try:
                header = next(rows)
            except StopIteration:
                break

            abbr_idx = None
            name_idx = None
            for i, cell in enumerate(header):
                val = str(cell.value).strip() if cell.value is not None else ""
                v = val.lower()
                if abbr_idx is None and any(k in v for k in ["缩写", "abbrev"]):
                    abbr_idx = i
                if name_idx is None and any(k in v for k in ["期刊名称", "journal", "title"]):
                    name_idx = i

            if abbr_idx is None or name_idx is None:
                break

            # 创建 / 清空 JCR 缩写临时表
            await db.execute(text("""
                CREATE TEMP TABLE IF NOT EXISTS _jcr_abbr(
                  full_name varchar(300),
                  nlm_abbr varchar(100)
                );
            """))
            await db.execute(text("TRUNCATE _jcr_abbr;"))

            for row in rows:
                abbr_raw = row[abbr_idx].value if abbr_idx < len(row) else None
                name_raw = row[name_idx].value if name_idx < len(row) else None
                if not abbr_raw or not name_raw:
                    continue
                abbr = str(abbr_raw).strip()
                full_name = str(name_raw).strip()
                if not abbr or not full_name:
                    continue
                await db.execute(
                    text("INSERT INTO _jcr_abbr(full_name, nlm_abbr) VALUES (:full_name, :abbr)"),
                    {"full_name": full_name, "abbr": abbr}
                )

            # 用 full_name 对齐 journal_metadata，补充 nlm_abbr
            await db.execute(text("""
                UPDATE journal_metadata jm
                SET nlm_abbr = a.nlm_abbr,
                    updated_at = NOW()
                FROM _jcr_abbr a
                WHERE jm.full_name = a.full_name
                  AND (jm.nlm_abbr IS NULL OR jm.nlm_abbr = '' OR jm.nlm_abbr = jm.issn);
            """))

            break  # 只处理第一个 JCR sheet

    return total_upserts


async def enrich_if_for_recent(db: AsyncSession, days: int = 90) -> int:
    """回填近期文献（默认90天内）的 IF5。

    逻辑：
    1. 优先使用 literature.journal_issn = journal_metadata.issn 精确匹配；
    2. 若 journal_issn 为空，则尝试使用 PubMed 期刊缩写（literature.journal_name = journal_metadata.nlm_abbr）；
    3. 同时将匹配到的 ISSN 回写到 literature.journal_issn，方便后续直接用 ISSN 匹配。
    返回受影响行数。
    """
    q = text(
        """
        UPDATE literature l
        SET journal_if_5y = jm.if_5y,
            journal_issn = COALESCE(l.journal_issn, jm.issn),
            journal_zone = COALESCE(l.journal_zone, jm.cas_zone)
        FROM journal_metadata jm
        WHERE (
                l.journal_issn = jm.issn
             OR (
                 l.journal_issn IS NULL
                 AND lower(trim(l.journal_name)) = lower(trim(jm.nlm_abbr))
             )
              )
          AND (l.journal_if_5y IS DISTINCT FROM jm.if_5y)
          AND l.publication_date >= (CURRENT_DATE - INTERVAL ':d days')
        """.replace(":d", str(days))
    )
    res = await db.execute(q)
    # rowcount 在不同驱动下表现不同，做个兜底
    try:
        return int(res.rowcount or 0)
    except Exception:
        return 0
